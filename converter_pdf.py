#!/usr/bin/env python3
"""Legacy UNIMED converter facade.

The parsing implementation lives in :mod:`rateiosrh.parsers.unimed`.  This
module keeps the historical CLI and public functions available to callers.
"""

from __future__ import annotations

import argparse
import logging
from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from rateiosrh.core.exceptions import PdfStructureError
from rateiosrh.core.pdf_reader import PdfDocument, TextSpan, clean, group_rows
from rateiosrh.parsers.unimed import (
    DEFAULT_LAYOUT,
    HEADERS,
    ExtractionResult,
    LayoutConfig,
    Record,
    TitularTotal,
    UnimedParser,
    _compose_operation,
    _field_text,
    _parse_decimal,
    _record_from_row,
    _title_parts,
    extract_records,
    split_code_and_name,
    validate_records,
)


LOGGER = logging.getLogger("converter_pdf")


def _style_sheet(sheet, widths: Sequence[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{sheet.max_row}"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.number_format = "@"


def write_workbook(result: ExtractionResult, output_path: Path) -> Path:
    """Write the main launch table and a separate validation-observation sheet."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lançamentos"
    sheet.append(HEADERS)
    for record in result.records:
        sheet.append(
            [
                record.codigo_beneficiario,
                record.beneficiario,
                record.op_lancamento,
                record.filial,
                record.centro_custo,
                record.quantidade,
                record.valor_unit,
                record.valor_bonif,
                record.valor_total,
                record.mensagem,
            ]
        )
    _style_sheet(sheet, [25, 38, 24, 10, 14, 13, 15, 16, 15, 28])

    observations = workbook.create_sheet("Observações")
    observations.append(["TIPO", "DETALHE"])
    for message in validate_records(result):
        observations.append(["VALIDAÇÃO", message])
    observations.column_dimensions["A"].width = 16
    observations.column_dimensions["B"].width = 120
    observations.freeze_panes = "A2"
    observations.auto_filter.ref = f"A1:B{observations.max_row}"
    for cell in observations[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for row in observations.iter_rows():
        for cell in row:
            cell.number_format = "@"

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    LOGGER.info("Excel gerado: %s", output_path)
    return output_path


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Converte um relatório analítico PDF do layout de referência para Excel."
    )
    parser.add_argument("pdf", type=Path, help="arquivo PDF de entrada")
    parser.add_argument("-o", "--output", type=Path, help="arquivo XLSX de saída")
    parser.add_argument("-v", "--verbose", action="store_true", help="habilita logs detalhados")
    return parser


def run(argv: Optional[Sequence[str]] = None) -> int:
    parser = build_argument_parser()
    args = parser.parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s: %(message)s",
    )

    if not args.pdf.is_file():
        LOGGER.error("Arquivo PDF não encontrado: %s", args.pdf)
        return 2
    output = args.output or args.pdf.with_suffix(".xlsx")
    try:
        result = extract_records(args.pdf)
        observations = validate_records(result)
        if observations:
            LOGGER.warning("Validações com observações: %d", len(observations))
            for message in observations:
                LOGGER.warning(message)
        unique_beneficiaries = {
            record.codigo_beneficiario
            for record in result.records
            if record.codigo_beneficiario and record.source_beneficiary_repeated
        }
        LOGGER.info("Beneficiários/dependentes identificados: %d", len(unique_beneficiaries))
        write_workbook(result, output)
        return 0
    except (FileNotFoundError, PdfStructureError, OSError, ValueError) as exc:
        LOGGER.error("Falha na conversão: %s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(run())

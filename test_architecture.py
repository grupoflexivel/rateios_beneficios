from pathlib import Path

import pandas as pd
import pytest

from rateiosrh.core.pdf_reader import PdfDocument, group_rows
from rateiosrh.core.exceptions import (
    DataFrameValidationError,
    LayoutMismatchError,
    UnknownParserError,
)
from rateiosrh.parsers.base import BaseParser
from rateiosrh.parsers.unimed import UnimedParser


class _ContractParser(BaseParser):
    parser_id = "contract"
    display_name = "Contract parser"
    columns = ["A", "B"]

    def validate_layout(self, pdf_path: Path) -> bool:
        return True

    def parse(self, pdf_path: Path) -> pd.DataFrame:
        return pd.DataFrame(columns=self.columns)


def test_application_packages_are_importable():
    import rateiosrh.core
    import rateiosrh.parsers
    import rateiosrh.services
    import rateiosrh.web

    assert rateiosrh.core is not None


def test_reader_finds_ten_pages_in_unimed_reference():
    document = PdfDocument(Path("ANALITICO_TAXA_0054041451-1.PDF").read_bytes())
    assert len(document.page_streams()) == 10


def test_reader_finds_seven_pages_in_bradesco_dental_reference():
    document = PdfDocument(Path("FLEXIV_1166046_1_082026.pdf").read_bytes())
    assert len(document.page_streams()) == 7


def test_common_reader_removes_inline_images_with_pdf_whitespace():
    import zlib

    from rateiosrh.core.pdf_reader import remove_inline_images

    image = zlib.compress(b"image payload")
    stream = (
        b"q BI \r\n/W 1\r\n/H 1\r\n/F /FlateDecode\r\nID\r\n"
        + image
        + b"\r\n EI\t\r\nBT\n(visible)Tj\nET"
    )

    cleaned = remove_inline_images(stream)

    assert image not in cleaned
    assert b"(visible)Tj" in cleaned


@pytest.mark.parametrize(
    ("pages_type", "page_type"),
    [
        (b"/Type /Pages", b"/Type /Page"),
        (b"/Type/Pages", b"/Type/Page"),
        (b"/Type /Pages", b"/Type/Page"),
        (b"/Type/Pages", b"/Type /Page"),
    ],
)
def test_reader_accepts_page_type_spacing_variants(pages_type, page_type):
    data = (
        b"1 0 obj\n<< "
        + pages_type
        + b" /Kids [2 0 R] >>\nendobj\n"
        + b"2 0 obj\n<< "
        + page_type
        + b" /Contents 3 0 R >>\nendobj\n"
        + b"3 0 obj\n<< /Length 5 >>\nstream\nBT ET\nendstream\nendobj\n"
    )

    assert PdfDocument(data).page_streams() == [b"BT ET"]


def test_group_rows_uses_requested_precision_without_pre_rounding_spans():
    stream = (
        b"1 0 0 1 0 10.001 Tm (first) Tj "
        b"1 0 0 1 0 10.004 Tm (second) Tj"
    )

    assert [y for y, _ in group_rows(stream, 3)] == [10.004, 10.001]


def test_base_parser_requires_parse_implementation():
    with pytest.raises(TypeError):
        BaseParser()


def test_schema_validation_accepts_exact_columns_without_inspecting_rows():
    from rateiosrh.services.validation import ValidationService

    frame = pd.DataFrame([[object(), object()]], columns=["A", "B"])

    assert ValidationService.validate_schema(frame, ["A", "B"]) == []


@pytest.mark.parametrize(
    "columns",
    [
        ["B", "A"],
        ["A"],
        ["A", "C"],
        ["A", "B", "C"],
    ],
)
def test_schema_validation_rejects_wrong_column_quantity_name_or_order(columns):
    from rateiosrh.services.validation import ValidationService

    frame = pd.DataFrame([list(range(len(columns)))], columns=columns)

    assert ValidationService.validate_schema(frame, ["A", "B"]) == [
        "Colunas ausentes ou fora de ordem."
    ]


def test_base_parser_dataframe_validation_does_not_mutate_dataframe():
    parser = _ContractParser()
    frame = pd.DataFrame([["x", "y"]], columns=["A", "B"])
    before = frame.copy(deep=True)

    assert parser.validate_dataframe(frame) == []
    pd.testing.assert_frame_equal(frame, before)


def test_base_parser_get_columns_returns_a_copy_of_declared_columns():
    parser = _ContractParser()

    columns = parser.get_columns()
    columns.append("C")

    assert columns == ["A", "B", "C"]
    assert parser.get_columns() == ["A", "B"]


def test_unimed_parser_returns_the_validated_dataframe():
    frame = UnimedParser().parse(Path("ANALITICO_TAXA_0054041451-1.PDF"))

    assert len(frame) == 185
    assert frame.columns.tolist() == [
        "CÓDIGO BENEFICIÁRIO",
        "BENEFICIÁRIO",
        "OP. LANÇAMENTO",
        "FILIAL",
        "CENTROCUSTO",
        "QUANT.",
        "VALOR UNIT.",
        "VALOR BONIF.",
        "VALOR TOTAL",
        "MENSAGEM",
    ]
    assert all(str(dtype) == "string" for dtype in frame.dtypes)
    assert frame.iloc[0]["VALOR TOTAL"] == "407,19"
    assert frame.iloc[0]["CÓDIGO BENEFICIÁRIO"] == "0976.0663.000060-00"


def test_unimed_parser_declares_metadata_and_rejects_other_layout():
    parser = UnimedParser()

    assert parser.parser_id == "unimed"
    assert parser.display_name == "Converter UNIMED"
    assert parser.validate_layout(Path("ANALITICO_TAXA_0054041451-1.PDF"))
    assert not parser.validate_layout(Path("FLEXIV_1166046_1_082026.pdf"))


def test_unimed_parser_dataframe_validation_reports_empty_required_field():
    parser = UnimedParser()
    frame = parser.parse(Path("ANALITICO_TAXA_0054041451-1.PDF"))
    frame.loc[0, "VALOR TOTAL"] = ""
    before = frame.copy(deep=True)

    assert parser.validate_dataframe(frame) == [
        "Lançamento 1: campo obrigatório vazio: valor total."
    ]
    pd.testing.assert_frame_equal(frame, before)


def test_registry_lists_all_implemented_converters():
    from rateiosrh.registry import get_parser, list_parsers

    assert [item["id"] for item in list_parsers()] == [
        "unimed",
        "bradesco_dental",
        "bradesco_seguros",
    ]
    assert get_parser("bradesco_dental").parser_id == "bradesco_dental"
    assert get_parser("bradesco_seguros").parser_id == "bradesco_seguros"


def test_registry_rejects_unknown_converter():
    from rateiosrh.registry import get_parser

    with pytest.raises(UnknownParserError):
        get_parser("bradesco_seguro")


def test_conversion_service_returns_dataframe():
    from rateiosrh.services.conversion import ConversionService

    output = ConversionService().convert(
        "bradesco_dental", Path("FLEXIV_1166046_1_082026.pdf")
    )

    assert output.parser_id == "bradesco_dental"
    assert output.dataframe.shape == (279, 13)
    assert output.page_count == 7
    assert isinstance(output.observations, list)


def test_conversion_service_preserves_parser_observations():
    from rateiosrh.services.conversion import ConversionService

    output = ConversionService().convert(
        "unimed", Path("ANALITICO_TAXA_0054041451-1.PDF")
    )

    assert len(output.observations) == 9
    assert "beneficiário não foi repetido" in output.observations[0]


def test_conversion_service_rejects_incompatible_layout():
    from rateiosrh.services.conversion import ConversionService

    with pytest.raises(LayoutMismatchError):
        ConversionService().convert(
            "bradesco_dental", Path("ANALITICO_TAXA_0054041451-1.PDF")
        )


def test_conversion_service_raises_dataframe_validation_error(monkeypatch):
    from rateiosrh.registry import get_parser
    from rateiosrh.services.conversion import ConversionService

    parser = get_parser("bradesco_dental")
    monkeypatch.setattr(parser, "validate_layout", lambda pdf_path: True)
    monkeypatch.setattr(parser, "parse", lambda pdf_path: pd.DataFrame())
    monkeypatch.setattr(
        "rateiosrh.services.conversion.get_parser", lambda parser_id: parser
    )

    with pytest.raises(DataFrameValidationError) as error:
        ConversionService().convert(
            "bradesco_dental", Path("FLEXIV_1166046_1_082026.pdf")
        )

    assert error.value.errors == ["Colunas ausentes ou fora de ordem."]


def test_excel_exporter_writes_exact_dataframe_columns(tmp_path):
    from openpyxl import load_workbook

    from rateiosrh.services.conversion import ConversionService
    from rateiosrh.services.excel_exporter import ExcelExporter

    output = tmp_path / "bradesco.xlsx"
    result = ConversionService().convert(
        "bradesco_dental", Path("FLEXIV_1166046_1_082026.pdf")
    )

    ExcelExporter.export(result.dataframe, output)

    workbook = load_workbook(output)
    sheet = workbook["Lançamentos"]
    assert output.exists()
    assert [cell.value for cell in sheet[1]] == result.dataframe.columns.tolist()
    assert sheet.max_row == 280
    assert sheet.auto_filter.ref == "A1:M280"
    assert sheet.freeze_panes == "A2"


def test_excel_exporter_writes_observations_without_mutating_dataframe(tmp_path):
    from openpyxl import load_workbook

    from rateiosrh.services.excel_exporter import ExcelExporter

    frame = pd.DataFrame(
        [["001", "15,00-", "=1+1"]], columns=["Código", "Valor", "Fórmula"]
    )
    before = frame.copy(deep=True)
    output = tmp_path / "observacoes.xlsx"

    ExcelExporter.export(frame, output, ["linha de teste"])

    workbook = load_workbook(output)
    assert "Observações" in workbook.sheetnames
    assert workbook["Observações"]["B2"].value == "linha de teste"
    assert workbook["Lançamentos"]["C2"].value == "=1+1"
    assert workbook["Lançamentos"]["C2"].data_type == "s"
    pd.testing.assert_frame_equal(frame, before)


def test_excel_exporter_does_not_leave_partial_output_on_save_error(
    tmp_path, monkeypatch
):
    from openpyxl import Workbook

    from rateiosrh.services.excel_exporter import ExcelExporter

    output = tmp_path / "parcial.xlsx"
    frame = pd.DataFrame([["x"]], columns=["Código"])

    def fail_save(workbook, filename):
        Path(filename).write_bytes(b"partial")
        raise OSError("falha simulada")

    monkeypatch.setattr(Workbook, "save", fail_save)

    with pytest.raises(OSError, match="falha simulada"):
        ExcelExporter.export(frame, output)

    assert not output.exists()
    assert list(tmp_path.glob(".parcial.xlsx.*.tmp")) == []

"""Generic Excel output for validated conversion DataFrames."""

from pathlib import Path
import os
import tempfile
from typing import Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Write a DataFrame to XLSX without applying layout-specific rules."""

    _MAIN_SHEET = "Lançamentos"
    _OBSERVATIONS_SHEET = "Observações"

    @staticmethod
    def _cell_value(value: object) -> object:
        if value is None or value is pd.NA:
            return None
        try:
            if bool(pd.isna(value)):
                return None
        except (TypeError, ValueError):
            pass
        return value

    @classmethod
    def _format_sheet(cls, sheet, width_count: int, fill_color: str) -> None:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(width_count)}{sheet.max_row}"
            if width_count
            else None
        )
        header_fill = PatternFill("solid", fgColor=fill_color)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for row in sheet.iter_rows():
            for cell in row:
                cell.number_format = "@"

    @classmethod
    def _set_widths(cls, sheet, width_count: int) -> None:
        for index in range(1, width_count + 1):
            values = [
                len(str(sheet.cell(row=row, column=index).value or ""))
                for row in range(1, sheet.max_row + 1)
            ]
            sheet.column_dimensions[get_column_letter(index)].width = min(
                max(max(values, default=10) + 2, 10), 60
            )

    @staticmethod
    def _append_textual_row(sheet, values: Sequence[object]) -> None:
        sheet.append(list(values))
        for cell in sheet[sheet.max_row]:
            if cell.value is not None:
                cell.value = str(cell.value)
                cell.data_type = "s"

    @classmethod
    def export(
        cls,
        df: pd.DataFrame,
        output_path: Path,
        observations: Sequence[str] = (),
    ) -> Path:
        """Write ``df`` and optional observations, returning the output path."""

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = cls._MAIN_SHEET
        columns = list(df.columns)
        cls._append_textual_row(sheet, columns)
        for row in df.itertuples(index=False, name=None):
            cls._append_textual_row(
                sheet, [cls._cell_value(value) for value in row]
            )
        cls._format_sheet(sheet, len(columns), "D9EAF7")
        cls._set_widths(sheet, len(columns))

        observation_values = list(observations)
        if observation_values:
            observation_sheet = workbook.create_sheet(cls._OBSERVATIONS_SHEET)
            cls._append_textual_row(observation_sheet, ["TIPO", "DETALHE"])
            for message in observation_values:
                cls._append_textual_row(observation_sheet, ["VALIDAÇÃO", message])
            cls._format_sheet(observation_sheet, 2, "FFF2CC")
            cls._set_widths(observation_sheet, 2)

        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary_path = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=output.parent,
                prefix=f".{output.name}.",
                suffix=".tmp",
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)
            workbook.save(temporary_path)
            os.replace(temporary_path, output)
            temporary_path = None
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
        return output
